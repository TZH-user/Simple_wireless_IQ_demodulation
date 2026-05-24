from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(r"D:\VSCodeCMAKEProject\Simple_wireless_IQ_demodulation")
OUT = ROOT / "analysis_reports" / "频谱分析报告_示例模板_修正版_v5.xlsx"

TESTS = [
    {
        "id": "T001",
        "log": Path(r"D:\串口调试助手-V3.1\logs\serial_export_20260520_160813.txt"),
        "lo_offset_hz": 5000,
        "note": "有线连接，信号源1.5Vpp，无直流偏置，120MHz CW，5k本振偏置",
    },
    {
        "id": "T002",
        "log": Path(r"D:\串口调试助手-V3.1\logs\serial_export_20260520_160643.txt"),
        "lo_offset_hz": 7000,
        "note": "有线连接，信号源1.5Vpp，无直流偏置，120MHz CW，7k本振偏置",
    },
    {
        "id": "T003",
        "log": Path(r"D:\串口调试助手-V3.1\logs\serial_export_20260520_160933.txt"),
        "lo_offset_hz": 10000,
        "note": "有线连接，信号源1.5Vpp，无直流偏置，120MHz CW，10k本振偏置",
    },
]

SPEC_ID_TO_NAME = {1: "IQ", 2: "包络", 3: "相位"}
SPEC_NAME_TO_SHEET = {"IQ": "IQ谱数据", "包络": "包络谱数据", "相位": "相位谱数据"}
SPEC_ORDER = {"IQ": 1, "包络": 2, "相位": 3}
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")


def parse_log(test: dict) -> dict:
    text = test["log"].read_text(encoding="utf-8", errors="ignore")
    center_match = re.search(r"sweep done center=(\d+)Hz", text)
    depths = [int(x) for x in re.findall(r"analyze: depth=(\d+)pm", text)]
    done_match = re.search(r"calculation complete center=(\d+)Hz mod=(\d+)Hz depth=(\d+)pm", text)

    records = []
    for freq, mag, phase, bin_no, spec_id in re.findall(
        r"spec:([0-9.]+),([0-9.]+),([-0-9.]+),(\d+),(\d+)", text
    ):
        sid = int(spec_id)
        if sid not in SPEC_ID_TO_NAME:
            continue
        records.append(
            {
                "测试编号": test["id"],
                "频谱类型": SPEC_ID_TO_NAME[sid],
                "频率_Hz": float(freq),
                "幅度": float(mag),
                "相位_rad": float(phase),
                "Bin": int(bin_no),
                "数据源文件": test["log"].name,
            }
        )

    return {
        **test,
        "center_hz": int(done_match.group(1)) if done_match else int(center_match.group(1)) if center_match else None,
        "depth_pm": int(done_match.group(3)) if done_match else round(sum(depths) / len(depths)) if depths else None,
        "records": records,
    }


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def set_widths(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def append_rows(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    style_header(ws)


def add_line_chart(ws_chart, ws_chart_data, title, start_row, end_row, pos, value_col, value_title, is_phase=False, log_y=False):
    chart = LineChart()
    chart.title = f"{title}（0~200kHz，{value_title}）"
    chart.style = 2
    chart.width = 36
    chart.height = 15
    chart.x_axis.title = "频率_Hz"
    chart.y_axis.title = value_title
    chart.y_axis.majorGridlines = None
    chart.x_axis.tickLblSkip = 50
    chart.legend.position = "r"
    if is_phase:
        chart.y_axis.scaling.min = -3.2
        chart.y_axis.scaling.max = 3.2
        chart.y_axis.majorUnit = 0.8
    if log_y:
        chart.y_axis.scaling.logBase = 10
        chart.y_axis.scaling.min = 1

    data = Reference(ws_chart_data, min_col=value_col, min_row=start_row, max_row=end_row)
    cats = Reference(ws_chart_data, min_col=3, min_row=start_row, max_row=end_row)
    chart.add_data(data, titles_from_data=False)
    if chart.series:
        chart.series[0].tx = SeriesLabel(v=value_title)
    chart.set_categories(cats)
    ws_chart.add_chart(chart, pos)


def main():
    parsed_tests = [parse_log(t) for t in TESTS]
    wb = Workbook()
    wb.remove(wb.active)

    ws_tests = wb.create_sheet("测试记录")
    test_headers = [
        "测试编号",
        "日期",
        "数据源文件",
        "载波频率_Hz",
        "调制类型",
        "调制速率_Hz",
        "调制参数",
        "连接方式",
        "信号源峰峰值_V",
        "信号源直流偏置_V",
        "本振偏置_Hz",
        "包络深度_pm",
        "备注",
    ]
    test_rows = []
    for t in parsed_tests:
        test_rows.append(
            {
                "测试编号": t["id"],
                "日期": "2026-05-20",
                "数据源文件": t["log"].name,
                "载波频率_Hz": t["center_hz"],
                "调制类型": "CW",
                "调制速率_Hz": "",
                "调制参数": "",
                "连接方式": "有线",
                "信号源峰峰值_V": 1.5,
                "信号源直流偏置_V": 0,
                "本振偏置_Hz": t["lo_offset_hz"],
                "包络深度_pm": t["depth_pm"],
                "备注": t["note"],
            }
        )
    test_rows.append(
        {
            "测试编号": "Txxx",
            "调制类型": "AM/ASK/FSK/PSK",
            "连接方式": "无线/有线",
            "备注": "模板行：后续新增测试时填基础条件",
        }
    )
    append_rows(ws_tests, test_headers, test_rows)
    set_widths(ws_tests, [12, 13, 34, 16, 12, 14, 14, 12, 17, 18, 14, 14, 56])

    spectrum_ranges = {}
    data_headers = ["测试编号", "频率_Hz", "幅度", "相位_rad", "Bin", "数据源文件"]
    for spec_name, sheet_name in SPEC_NAME_TO_SHEET.items():
        ws = wb.create_sheet(sheet_name)
        rows = [r for t in parsed_tests for r in t["records"] if r["频谱类型"] == spec_name]
        append_rows(ws, data_headers, rows)
        set_widths(ws, [12, 14, 14, 14, 10, 34])
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=4):
            row[0].number_format = "0.000"
            row[1].number_format = "0.000"
            row[2].number_format = "0.000"

        row_cursor = 2
        spectrum_ranges[spec_name] = {}
        for t in parsed_tests:
            count = sum(1 for r in t["records"] if r["频谱类型"] == spec_name)
            if count > 0:
                spectrum_ranges[spec_name][t["id"]] = (row_cursor, row_cursor + count - 1)
                row_cursor += count

    ws_long = wb.create_sheet("频谱长表")
    long_headers = ["测试编号", "频率_Hz", "频谱类型", "幅度", "相位_rad", "Bin", "数据源文件"]
    long_rows = []
    for t in parsed_tests:
        long_rows.extend(t["records"])
    long_rows.sort(key=lambda r: (r["测试编号"], r["频率_Hz"], SPEC_ORDER[r["频谱类型"]], r["Bin"]))
    append_rows(ws_long, long_headers, long_rows)
    set_widths(ws_long, [12, 14, 12, 14, 14, 10, 34])
    for row in ws_long.iter_rows(min_row=2, max_row=ws_long.max_row, min_col=2, max_col=5):
        row[0].number_format = "0.000"
        row[2].number_format = "0.000"
        row[3].number_format = "0.000"

    ws_chart_data = wb.create_sheet("图表数据")
    chart_data_ranges = {}
    chart_headers = ["测试编号", "频谱类型", "频率_Hz", "幅度", "相位_rad", "Bin"]
    ws_chart_data.append(chart_headers)
    style_header(ws_chart_data)
    set_widths(ws_chart_data, [12, 12, 14, 14, 14, 10])
    row_cursor = 2
    for t in parsed_tests:
        chart_data_ranges[t["id"]] = {}
        for spec_name in ["IQ", "包络", "相位"]:
            rows = [
                r for r in t["records"]
                if (r["频谱类型"] == spec_name and r["频率_Hz"] <= 200000)
            ]
            if not rows:
                continue
            start_row = row_cursor
            for r in rows:
                ws_chart_data.append([
                    r["测试编号"],
                    r["频谱类型"],
                    r["频率_Hz"],
                    r["幅度"],
                    r["相位_rad"],
                    r["Bin"],
                ])
                row_cursor += 1
            end_row = row_cursor - 1
            chart_data_ranges[t["id"]][spec_name] = (start_row, end_row)
    ws_chart_data.freeze_panes = "A2"
    ws_chart_data.auto_filter.ref = ws_chart_data.dimensions

    ws_count = wb.create_sheet("谱点统计")
    ws_count.append(["测试编号", "IQ点数", "包络点数", "相位点数", "说明"])
    for t in parsed_tests:
        ws_count.append(
            [
                t["id"],
                sum(1 for r in t["records"] if r["频谱类型"] == "IQ"),
                sum(1 for r in t["records"] if r["频谱类型"] == "包络"),
                sum(1 for r in t["records"] if r["频谱类型"] == "相位"),
                "IQ 为复数谱 4096 点；包络/相位为 RFFT 半谱 2048 点",
            ]
        )
    style_header(ws_count)
    set_widths(ws_count, [12, 12, 12, 12, 56])

    for t in parsed_tests:
        ws_chart = wb.create_sheet(f"{t['id']}_图表")
        ws_chart.sheet_view.zoomScale = 75
        ws_chart["A1"] = f"{t['id']} 图表页：默认显示 0~200kHz，完整频谱数据见各谱数据页。"
        ws_chart["A1"].font = Font(bold=True, size=13)
        ws_chart["A1"].fill = NOTE_FILL
        ws_chart["A2"] = f"条件：{t['note']}；包络深度={t['depth_pm']}pm"
        ws_chart["A2"].fill = NOTE_FILL
        set_widths(ws_chart, [22] * 24)

        chart_defs = [
            ("IQ", "A4", "Q4", "A42"),
            ("包络", "A80", "Q80", "A118"),
            ("相位", "A156", "Q156", "A194"),
        ]
        for spec_name, amp_pos, log_amp_pos, phase_pos in chart_defs:
            if spec_name not in chart_data_ranges[t["id"]]:
                continue
            start_row, end_row = chart_data_ranges[t["id"]][spec_name]
            add_line_chart(
                ws_chart,
                ws_chart_data,
                f"{t['id']} {spec_name}谱幅度",
                start_row,
                end_row,
                amp_pos,
                value_col=4,
                value_title="幅度",
            )
            add_line_chart(
                ws_chart,
                ws_chart_data,
                f"{t['id']} {spec_name}谱幅度对数坐标",
                start_row,
                end_row,
                log_amp_pos,
                value_col=4,
                value_title="幅度_log10",
                log_y=True,
            )
            add_line_chart(
                ws_chart,
                ws_chart_data,
                f"{t['id']} {spec_name}谱相位",
                start_row,
                end_row,
                phase_pos,
                value_col=5,
                value_title="相位_rad",
                is_phase=True,
            )

    ws_note = wb.create_sheet("模板说明")
    notes = [
        ["项目", "说明"],
        ["测试记录", "每次实验只占一行，保存实验条件和人工补充信息，不存放逐频点数据。"],
        ["各谱数据页", "每个频点一行；同一测试编号会重复多行，这是后续追加测试和筛选的正确形式。"],
        ["频谱长表", "三类谱按 测试编号 + 频率 + 谱类型 交错排列，便于透视表和跨谱对比。"],
        ["图表页", "每次测试固定三类谱；每类谱左侧为线性幅度图，右侧为幅度对数坐标图，下方为相位图；图表使用专用 0~200kHz 数据区，完整数据不裁剪，保存在数据页。"],
        ["缺失谱", "若后续日志未输出某类谱，该谱数据页保持空白，对应图表不生成。"],
    ]
    for row in notes:
        ws_note.append(row)
    style_header(ws_note)
    set_widths(ws_note, [18, 88])
    for row in ws_note.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT)
    print(f"WROTE: {OUT}")
    print(f"LONG_ROWS: {len(long_rows)}")
    print(
        "COUNTS: "
        + ", ".join(
            f"{t['id']} IQ={sum(1 for r in t['records'] if r['频谱类型'] == 'IQ')} "
            f"ENV={sum(1 for r in t['records'] if r['频谱类型'] == '包络')} "
            f"PHASE={sum(1 for r in t['records'] if r['频谱类型'] == '相位')}"
            for t in parsed_tests
        )
    )


if __name__ == "__main__":
    main()
