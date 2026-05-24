from __future__ import annotations

import argparse
import re
from pathlib import Path
from tkinter import Tk, filedialog

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SPEC_ID_TO_NAME = {
    1: "IQ",
    2: "包络",
    3: "相位",
}

SPEC_ORDER = {
    "IQ": 1,
    "包络": 2,
    "相位": 3,
}

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
CHART_MAX_HZ = 200000


def parse_log(log_path: Path, test_id: str) -> dict[str, object]:
    """解析串口频谱日志，返回测试元数据和逐频点数据。"""
    text = log_path.read_text(encoding="utf-8", errors="ignore")
    center_match = re.search(r"sweep done center=(\d+)Hz", text)
    done_match = re.search(r"calculation complete center=(\d+)Hz mod=(\d+)Hz depth=(\d+)pm", text)
    depths = [int(x) for x in re.findall(r"analyze: depth=(\d+)pm", text)]

    center_hz = ""
    mod_hz = ""
    depth_pm = ""
    if done_match:
        center_hz = done_match.group(1)
        mod_hz = done_match.group(2)
        depth_pm = done_match.group(3)
    elif center_match:
        center_hz = center_match.group(1)
        if depths:
            depth_pm = str(round(sum(depths) / len(depths)))

    records: list[dict[str, object]] = []
    for freq_hz, mag, phase_rad, bin_no, spec_id in re.findall(
        r"spec:([0-9.]+),([0-9.]+),([-0-9.]+),(\d+),(\d+)", text
    ):
        sid = int(spec_id)
        spec_name = SPEC_ID_TO_NAME.get(sid)
        if spec_name is None:
            continue
        records.append(
            {
                "测试编号": test_id,
                "频谱类型": spec_name,
                "频率_Hz": float(freq_hz),
                "幅度": float(mag),
                "相位_rad": float(phase_rad),
                "Bin": int(bin_no),
                "数据源文件": log_path.name,
            }
        )

    return {
        "测试编号": test_id,
        "数据源文件": log_path.name,
        "中心频率_Hz": center_hz,
        "识别调制频率_Hz": mod_hz,
        "包络深度_pm": depth_pm,
        "records": records,
    }


def split_log_paths(raw_paths: str) -> list[Path]:
    """允许用英文分号分隔多个日志路径。"""
    return [Path(item.strip().strip('"')) for item in raw_paths.split(";") if item.strip()]


def choose_files_with_dialog(args: argparse.Namespace) -> argparse.Namespace:
    """未传参数时使用文件窗口，避免手动复制长路径。"""
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    if not args.logs:
        log_paths = filedialog.askopenfilenames(
            title="选择串口导出日志文件，可多选",
            filetypes=[("串口日志", "*.txt *.log"), ("所有文件", "*.*")],
        )
        if log_paths:
            args.logs = ";".join(log_paths)

    if not args.out:
        out_path = filedialog.asksaveasfilename(
            title="选择输出 XLSX 文件名",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")],
        )
        if out_path:
            args.out = out_path

    root.destroy()
    return args


def ask_if_missing(args: argparse.Namespace) -> argparse.Namespace:
    args = choose_files_with_dialog(args)
    if not args.logs:
        args.logs = input("请输入串口日志路径，多个文件用英文分号分隔：").strip()
    if not args.out:
        args.out = input("请输入输出 XLSX 文件路径或文件名：").strip()
    return args


def style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def set_widths(ws, widths: list[float]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def append_table(ws, headers: list[str], rows: list[dict[str, object]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    style_header(ws)


def add_line_chart(
    ws_chart,
    ws_data,
    title: str,
    start_row: int,
    end_row: int,
    pos: str,
    value_col: int,
    value_title: str,
    is_phase: bool = False,
    log_y: bool = False,
) -> None:
    chart = LineChart()
    chart.title = title
    chart.style = 2
    chart.width = 36
    chart.height = 15
    chart.x_axis.title = "频率_Hz"
    chart.y_axis.title = value_title
    chart.y_axis.majorGridlines = None
    chart.x_axis.tickLblSkip = 50
    chart.legend.position = "r"

    if is_phase:
        chart.y_axis.scaling.min = -3.2
        chart.y_axis.scaling.max = 3.2
        chart.y_axis.majorUnit = 0.8
    if log_y:
        chart.y_axis.scaling.logBase = 10
        chart.y_axis.scaling.min = 1

    data = Reference(ws_data, min_col=value_col, min_row=start_row, max_row=end_row)
    cats = Reference(ws_data, min_col=3, min_row=start_row, max_row=end_row)
    chart.add_data(data, titles_from_data=False)
    if chart.series:
        chart.series[0].tx = SeriesLabel(v=value_title)
    chart.set_categories(cats)
    ws_chart.add_chart(chart, pos)


def build_workbook(parsed_tests: list[dict[str, object]], out_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    test_headers = ["测试编号", "数据源文件", "中心频率_Hz", "识别调制频率_Hz", "包络深度_pm", "备注"]
    test_rows = [
        {
            "测试编号": item["测试编号"],
            "数据源文件": item["数据源文件"],
            "中心频率_Hz": item["中心频率_Hz"],
            "识别调制频率_Hz": item["识别调制频率_Hz"],
            "包络深度_pm": item["包络深度_pm"],
            "备注": "自动转换生成，可手动补充调制类型、连接方式、信号源幅度等条件",
        }
        for item in parsed_tests
    ]
    ws_tests = wb.create_sheet("测试记录")
    append_table(ws_tests, test_headers, test_rows)
    set_widths(ws_tests, [12, 34, 16, 18, 14, 62])

    data_headers = ["测试编号", "频率_Hz", "幅度", "相位_rad", "Bin", "数据源文件"]
    all_records: list[dict[str, object]] = []
    for item in parsed_tests:
        all_records.extend(item["records"])

    for spec_name, sheet_name in [("IQ", "IQ谱数据"), ("包络", "包络谱数据"), ("相位", "相位谱数据")]:
        ws = wb.create_sheet(sheet_name)
        rows = [r for r in all_records if r["频谱类型"] == spec_name]
        append_table(ws, data_headers, rows)
        set_widths(ws, [12, 14, 14, 14, 10, 34])
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=4):
            for cell in row:
                cell.number_format = "0.000"

    ws_long = wb.create_sheet("频谱长表")
    long_headers = ["测试编号", "频率_Hz", "频谱类型", "幅度", "相位_rad", "Bin", "数据源文件"]
    long_rows = sorted(all_records, key=lambda r: (r["测试编号"], r["频率_Hz"], SPEC_ORDER[r["频谱类型"]], r["Bin"]))
    append_table(ws_long, long_headers, long_rows)
    set_widths(ws_long, [12, 14, 12, 14, 14, 10, 34])

    ws_chart_data = wb.create_sheet("图表数据")
    chart_headers = ["测试编号", "频谱类型", "频率_Hz", "幅度", "相位_rad", "Bin"]
    ws_chart_data.append(chart_headers)
    style_header(ws_chart_data)
    set_widths(ws_chart_data, [12, 12, 14, 14, 14, 10])
    chart_ranges: dict[str, dict[str, tuple[int, int]]] = {}
    row_cursor = 2
    for item in parsed_tests:
        test_id = str(item["测试编号"])
        chart_ranges[test_id] = {}
        for spec_name in ["IQ", "包络", "相位"]:
            rows = [r for r in item["records"] if r["频谱类型"] == spec_name and r["频率_Hz"] <= CHART_MAX_HZ]
            if not rows:
                continue
            start_row = row_cursor
            for r in rows:
                ws_chart_data.append([r["测试编号"], r["频谱类型"], r["频率_Hz"], r["幅度"], r["相位_rad"], r["Bin"]])
                row_cursor += 1
            chart_ranges[test_id][spec_name] = (start_row, row_cursor - 1)
    ws_chart_data.freeze_panes = "A2"
    ws_chart_data.auto_filter.ref = ws_chart_data.dimensions

    ws_count = wb.create_sheet("谱点统计")
    ws_count.append(["测试编号", "IQ点数", "包络点数", "相位点数", "说明"])
    for item in parsed_tests:
        test_id = item["测试编号"]
        records = item["records"]
        ws_count.append([
            test_id,
            sum(1 for r in records if r["频谱类型"] == "IQ"),
            sum(1 for r in records if r["频谱类型"] == "包络"),
            sum(1 for r in records if r["频谱类型"] == "相位"),
            "完整日志通常为 IQ 4096 点，包络 2048 点，相位 2048 点",
        ])
    style_header(ws_count)
    set_widths(ws_count, [12, 12, 12, 12, 58])

    for item in parsed_tests:
        test_id = str(item["测试编号"])
        ws_chart = wb.create_sheet(f"{test_id}_图表")
        ws_chart.sheet_view.zoomScale = 70
        ws_chart["A1"] = f"{test_id} 图表页：默认显示 0~200kHz，完整频谱数据见各谱数据页。"
        ws_chart["A1"].font = Font(bold=True, size=13)
        ws_chart["A1"].fill = NOTE_FILL
        ws_chart["A2"] = f"数据源：{item['数据源文件']}；中心频率={item['中心频率_Hz']}Hz；包络深度={item['包络深度_pm']}pm"
        ws_chart["A2"].fill = NOTE_FILL
        set_widths(ws_chart, [22] * 32)

        for spec_name, amp_pos, log_pos, phase_pos in [
            ("IQ", "A4", "Q4", "A42"),
            ("包络", "A80", "Q80", "A118"),
            ("相位", "A156", "Q156", "A194"),
        ]:
            if spec_name not in chart_ranges[test_id]:
                continue
            start_row, end_row = chart_ranges[test_id][spec_name]
            add_line_chart(ws_chart, ws_chart_data, f"{test_id} {spec_name}谱线性幅度", start_row, end_row, amp_pos, 4, "幅度")
            add_line_chart(ws_chart, ws_chart_data, f"{test_id} {spec_name}谱对数幅度", start_row, end_row, log_pos, 4, "幅度_log10", log_y=True)
            add_line_chart(ws_chart, ws_chart_data, f"{test_id} {spec_name}谱相位", start_row, end_row, phase_pos, 5, "相位_rad", is_phase=True)

    ws_note = wb.create_sheet("模板说明")
    for row in [
        ["项目", "说明"],
        ["测试记录", "每次实验一行，只放实验条件和摘要。"],
        ["频谱长表", "逐频点长表，三类谱按测试编号、频率、频谱类型交错排列。"],
        ["图表数据", "只截取 0~200kHz 供图表使用，完整数据仍在三类谱数据页。"],
        ["图表页", "每类谱包含线性幅度、对数幅度、相位三张图。"],
    ]:
        ws_note.append(row)
    style_header(ws_note)
    set_widths(ws_note, [18, 88])
    for row in ws_note.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="将串口频谱日志转换为 XLSX 频谱报告。")
    parser.add_argument("--logs", help="日志路径；多个日志用英文分号分隔。")
    parser.add_argument("--out", help="输出 XLSX 文件路径。")
    parser.add_argument("--test-prefix", default="T", help="测试编号前缀，默认 T。")
    args = ask_if_missing(parser.parse_args())

    log_paths = split_log_paths(args.logs)
    if not log_paths:
        raise SystemExit("未提供有效日志路径。")

    out_path = Path(args.out.strip().strip('"'))
    if out_path.suffix.lower() != ".xlsx":
        out_path = out_path.with_suffix(".xlsx")

    parsed_tests = []
    for index, log_path in enumerate(log_paths, start=1):
        if not log_path.exists():
            raise SystemExit(f"日志文件不存在：{log_path}")
        parsed_tests.append(parse_log(log_path, f"{args.test_prefix}{index:03d}"))

    build_workbook(parsed_tests, out_path)
    point_count = sum(len(item["records"]) for item in parsed_tests)
    print(f"XLSX 已生成：{out_path}")
    print(f"写入频谱点数：{point_count}")


if __name__ == "__main__":
    main()
